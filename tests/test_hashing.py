"""Covers app.hashing — the change-detection hash must be stable across
re-reads of unchanged data, and change when a cell value actually changes.
Always run against tmp_path copies, never the real workbook."""
import json

import openpyxl

from app.hashing import compute_source_hash
from tests.conftest import build_test_workbook


def test_hash_stable_across_repeated_reads(tmp_path):
    wb_path = build_test_workbook(tmp_path / "wb.xlsx")
    manual_path = tmp_path / "manual_inputs.json"
    manual_path.write_text(json.dumps({"program_status_override": None}))

    h1 = compute_source_hash(wb_path, manual_path)
    h2 = compute_source_hash(wb_path, manual_path)
    assert h1 == h2


def test_hash_changes_when_a_task_status_changes(tmp_path):
    wb_path = build_test_workbook(tmp_path / "wb.xlsx")
    manual_path = tmp_path / "manual_inputs.json"
    manual_path.write_text(json.dumps({"program_status_override": None}))

    before = compute_source_hash(wb_path, manual_path)

    wb = openpyxl.load_workbook(wb_path)
    ws = wb["Master Plan"]
    ws.cell(row=7, column=14, value="Complete")  # T-004's Status column
    wb.save(wb_path)

    after = compute_source_hash(wb_path, manual_path)
    assert before != after


def test_hash_changes_when_manual_inputs_change(tmp_path):
    wb_path = build_test_workbook(tmp_path / "wb.xlsx")
    manual_path = tmp_path / "manual_inputs.json"
    manual_path.write_text(json.dumps({"program_status_override": None}))
    before = compute_source_hash(wb_path, manual_path)

    manual_path.write_text(json.dumps({"program_status_override": "RED"}))
    after = compute_source_hash(wb_path, manual_path)
    assert before != after


def test_hash_ignores_the_read_me_sheet(tmp_path):
    """The Read Me sheet is instructional text, not tracked data — editing
    it must not be treated as a reportable change."""
    wb_path = build_test_workbook(tmp_path / "wb.xlsx")
    manual_path = tmp_path / "manual_inputs.json"
    manual_path.write_text(json.dumps({"x": 1}))

    wb = openpyxl.load_workbook(wb_path)
    wb.create_sheet("Read Me")
    wb["Read Me"]["A1"] = "Original instructions"
    wb.save(wb_path)
    before = compute_source_hash(wb_path, manual_path)

    wb2 = openpyxl.load_workbook(wb_path)
    wb2["Read Me"]["A1"] = "Completely different instructions"
    wb2.save(wb_path)
    after = compute_source_hash(wb_path, manual_path)

    assert before == after
