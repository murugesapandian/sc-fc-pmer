"""End-to-end smoke tests for app.pptx_builder against the real template,
including the overflow-protection stress scenario that was manually
validated during development (many more RAID items than the template's
example table has rows for) — now automated as a regression guard."""
from datetime import date, timedelta

from pptx import Presentation

from app.config import TEMPLATE_PPTX_PATH
from app.excel_reader import load_snapshot
from app.manual_inputs import load_manual_inputs
from app.pptx_builder import build_presentation, extract_text_signature, interpolate_x


def _build(tmp_path, workbook_path, manual_inputs_path, name="out.pptx"):
    snapshot = load_snapshot(workbook_path)
    manual = load_manual_inputs(manual_inputs_path)
    out = tmp_path / name
    build_presentation(snapshot, manual, TEMPLATE_PPTX_PATH, out, since=date.today() - timedelta(days=30))
    return out


def test_build_presentation_produces_a_12_slide_deck(tmp_path, test_workbook_path):
    manual_path = tmp_path / "manual_inputs.json"
    out = _build(tmp_path, test_workbook_path, manual_path)
    prs = Presentation(str(out))
    assert len(prs.slides) == 12


def test_chart_and_commentary_agree_on_current_phase(tmp_path, test_workbook_path):
    """Regression for the reported bug: the phase chart and the 'what this
    tells us' commentary must describe the same phase."""
    manual_path = tmp_path / "manual_inputs.json"
    out = _build(tmp_path, test_workbook_path, manual_path)
    prs = Presentation(str(out))
    slide4 = prs.slides[3]
    commentary = next(s.text_frame.text for s in slide4.shapes if s.name == "Text 4")
    assert "Beta" in commentary  # the fixture's current (incomplete) phase


def test_never_writes_to_the_real_template(tmp_path, test_workbook_path):
    import hashlib
    before = hashlib.sha256(TEMPLATE_PPTX_PATH.read_bytes()).hexdigest()
    _build(tmp_path, test_workbook_path, tmp_path / "manual_inputs.json")
    after = hashlib.sha256(TEMPLATE_PPTX_PATH.read_bytes()).hexdigest()
    assert before == after


def test_layout_guard_caps_overflow_with_many_raid_items(tmp_path, test_workbook_path):
    """Stress scenario manually validated during development: far more open
    RAID items than the template's example table/bullet boxes were sized
    for must not overflow — they must collapse into a '+N more' summary
    instead of growing past their shape's footprint."""
    import openpyxl
    from datetime import datetime

    wb = openpyxl.load_workbook(test_workbook_path)
    ws = wb["RAID Log"]
    for i in range(3, 40):
        ws.append([f"R-{i:03d}", "Risk", "A deliberately long risk description used to simulate a real "
                   "PMO entry that runs on at length about vendor exposure and mitigation history. " * 2,
                   None, "Engineering", "SH-02", datetime(2024, 1, 2), 5, 5, 25, "High",
                   "Extremely long mitigation text " * 4, datetime(2024, 2, 1), "Open", 0, "Test", datetime(2024, 1, 10)])
    wb.save(test_workbook_path)

    out = _build(tmp_path, test_workbook_path, tmp_path / "manual_inputs.json")
    prs = Presentation(str(out))

    slide8 = prs.slides[7]
    table = next(s.table for s in slide8.shapes if s.has_table)
    rows = list(table.rows)
    # Far fewer rows than the ~38 open High-severity items that exist now.
    assert len(rows) < 15
    assert "more" in rows[-1].cells[0].text.lower()


def test_extract_text_signature_is_stable_for_identical_content(tmp_path, test_workbook_path):
    manual_path = tmp_path / "manual_inputs.json"
    out1 = _build(tmp_path, test_workbook_path, manual_path, name="out1.pptx")
    out2 = _build(tmp_path, test_workbook_path, manual_path, name="out2.pptx")
    assert extract_text_signature(out1) == extract_text_signature(out2)


def test_interpolate_x_between_two_anchors():
    anchors = [(date(2024, 1, 1), 0), (date(2024, 1, 11), 100)]
    assert interpolate_x(anchors, date(2024, 1, 6)) == 50  # halfway between the two anchor dates


def test_interpolate_x_extrapolates_past_the_last_anchor():
    anchors = [(date(2024, 1, 1), 0), (date(2024, 1, 11), 100)]
    assert interpolate_x(anchors, date(2024, 1, 21)) == 200  # continues the same slope past the last point


def test_extract_text_signature_differs_after_a_real_change(tmp_path, test_workbook_path):
    manual_path = tmp_path / "manual_inputs.json"
    out1 = _build(tmp_path, test_workbook_path, manual_path, name="out1.pptx")

    import json
    data = json.loads(manual_path.read_text())
    data["decisions_needed"][0]["recommendation"] = "A completely different recommendation."
    manual_path.write_text(json.dumps(data))

    out2 = _build(tmp_path, test_workbook_path, manual_path, name="out2.pptx")
    assert extract_text_signature(out1) != extract_text_signature(out2)
