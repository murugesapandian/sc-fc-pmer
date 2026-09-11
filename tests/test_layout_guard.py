"""Covers app.layout_guard — the geometry-aware overflow protection.
Regression coverage for the specific bug found during development: line
capping must account for word-wrap (a long bullet costs several visual
lines, not one), or it silently lets far too much text through."""
from pptx import Presentation
from pptx.util import Inches, Pt

from app import layout_guard as lg


def _make_slide_with_boxes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    upper = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    lower = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))  # sits below, same x-range
    return slide, upper, lower


def test_available_height_below_uses_the_gap_to_the_next_shape():
    slide, upper, lower = _make_slide_with_boxes()
    available = lg.available_height_below(slide, upper)
    # Gap runs from upper.top (1in) to lower.top (4in) minus the margin —
    # comfortably more than upper's own native height (2in).
    assert available > upper.height


def test_available_height_below_falls_back_to_own_height_with_no_sibling():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    lone = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    assert lg.available_height_below(slide, lone) == lone.height


def test_available_height_below_ignores_shapes_with_no_x_overlap():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    # Positioned below but far to the right — shouldn't act as a floor.
    slide.shapes.add_textbox(Inches(8), Inches(2), Inches(2), Inches(1))
    assert lg.available_height_below(slide, box) == box.height


def test_cap_lines_to_fit_lets_short_lines_through_untouched():
    slide, box, _ = _make_slide_with_boxes()
    box.text_frame.paragraphs[0].add_run().text = "x"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    lines = ["short one", "short two", "short three"]
    result = lg.cap_lines_to_fit(box, lines, lg.available_height_below(slide, box))
    assert result == lines


def test_cap_lines_to_fit_truncates_many_long_wrapping_lines():
    """Regression: a first version of this function counted paragraphs, not
    wrapped visual lines, and let ~16 long bullets through a box sized for
    2-3. A long bullet must cost multiple 'lines' in the budget."""
    slide, box, _ = _make_slide_with_boxes()
    box.text_frame.paragraphs[0].add_run().text = "x"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    long_line = "This is a deliberately long line of text meant to simulate a wrapped multi-line bullet point in a narrow box. " * 2
    lines = [long_line] * 20
    available = box.height  # constrain to the box's own modest native height
    result = lg.cap_lines_to_fit(box, lines, available)
    assert len(result) < 20
    assert result[-1].startswith("+")  # the "+N more" summary line


def test_cap_rows_adds_summary_row_when_truncated():
    rows = [[f"row-{i}"] for i in range(10)]
    result = lg.cap_rows(rows, max_rows=3, note_row_factory=lambda n: [f"+{n} more"])
    assert len(result) == 3
    assert result[-1] == ["+8 more"]


def test_cap_rows_passthrough_when_it_fits():
    rows = [["a"], ["b"]]
    result = lg.cap_rows(rows, max_rows=5, note_row_factory=lambda n: [f"+{n}"])
    assert result == rows


def test_truncate_chars_short_text_unchanged():
    assert lg.truncate_chars("short", 100) == "short"


def test_truncate_chars_long_text_gets_ellipsis():
    result = lg.truncate_chars("a" * 300, 50)
    assert len(result) == 50
    assert result.endswith("…")
